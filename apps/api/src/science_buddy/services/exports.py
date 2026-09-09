import re
from collections.abc import Sequence
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def research_result_markdown(question: str, result: dict[str, Any]) -> str:
    lines = [
        "# Science Buddy Evidence Report",
        "",
        f"**Research question:** {question}",
        "",
        "> AI-assisted research output. Not for diagnosis, prescribing, or clinical decisions.",
        "",
        "## Answer",
        "",
        str(result.get("answer", "")),
        "",
        "## Verified claims and evidence",
        "",
    ]
    for index, claim in enumerate(result.get("claims", []), start=1):
        lines.extend([f"### Claim {index}", "", str(claim.get("statement", "")), ""])
        for evidence in claim.get("evidence", []):
            locator = evidence.get("source_locator", {})
            identifier = (
                locator.get("pmid")
                or locator.get("pmcid")
                or locator.get("doi")
                or "local"
            )
            lines.extend(
                [
                    f"- **{identifier}** — {locator.get('section_path', 'unknown section')}",
                    f"  - {evidence.get('text', '')}",
                ]
            )
        lines.append("")
    if result.get("gaps"):
        lines.extend(["## Evidence gaps", ""])
        lines.extend(f"- {value}" for value in result["gaps"])
        lines.append("")
    if result.get("conflicts"):
        lines.extend(["## Conflicts", ""])
        lines.extend(f"- {value}" for value in result["conflicts"])
    return "\n".join(lines).strip() + "\n"


def research_result_docx(question: str, result: dict[str, Any]) -> bytes:
    document = Document()
    document.add_heading("Science Buddy Evidence Report", level=0)
    document.add_paragraph(f"Research question: {question}")
    document.add_paragraph(
        "AI-assisted research output. Not for diagnosis, prescribing, or clinical decisions."
    )
    document.add_heading("Answer", level=1)
    document.add_paragraph(str(result.get("answer", "")))
    document.add_heading("Verified claims and evidence", level=1)
    for index, claim in enumerate(result.get("claims", []), start=1):
        document.add_heading(f"Claim {index}", level=2)
        document.add_paragraph(str(claim.get("statement", "")))
        for evidence in claim.get("evidence", []):
            locator = evidence.get("source_locator", {})
            identifier = (
                locator.get("pmid")
                or locator.get("pmcid")
                or locator.get("doi")
                or "local"
            )
            document.add_paragraph(
                f"{identifier} — {locator.get('section_path', 'unknown section')}",
                style="List Bullet",
            )
            document.add_paragraph(str(evidence.get("text", "")))
    if result.get("gaps"):
        document.add_heading("Evidence gaps", level=1)
        for value in result["gaps"]:
            document.add_paragraph(str(value), style="List Bullet")
    if result.get("conflicts"):
        document.add_heading("Conflicts", level=1)
        for value in result["conflicts"]:
            document.add_paragraph(str(value), style="List Bullet")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def synthesis_markdown(title: str, manuscript: str) -> bytes:
    value = manuscript.strip()
    if not value.startswith("# "):
        value = f"# {title}\n\n{value}"
    return (value.rstrip() + "\n").encode("utf-8")


def _paper_records(papers: Sequence[dict[str, Any]]) -> list[dict[str, str]]:
    """Normalize paper metadata dicts for citation-manager exports."""
    records: list[dict[str, str]] = []
    for paper in papers:
        authors = paper.get("authors") or []
        author_lines = []
        for author in authors:
            if not isinstance(author, dict):
                continue
            family = str(author.get("family_name") or author.get("full_name") or "")
            given = str(author.get("given_name") or "")
            if not family:
                continue
            author_lines.append(f"{family}, {given}".strip(" ,"))
        records.append(
            {
                "pmid": str(paper.get("pmid") or ""),
                "doi": str(paper.get("doi") or ""),
                "title": str(paper.get("title") or ""),
                "journal": str(paper.get("journal") or ""),
                "year": str(paper.get("publication_year") or ""),
                "authors": "; ".join(author_lines),
            }
        )
    return records


def _bibtex_key(record: dict[str, str]) -> str:
    if record["pmid"]:
        return f"sb_{record['pmid']}"
    source = record["doi"] or record["title"]
    slug = re.sub(r"[^a-zA-Z0-9]+", "", source)[:24]
    return f"sb_{slug}" if slug else "sb_reference"


def research_result_ris(
    question: str, result: dict[str, Any], papers: Sequence[dict[str, Any]]
) -> str:
    """Render the papers backing a research run as RIS records."""
    lines = ["# Science Buddy RIS export", "", f"# Research question: {question}", ""]
    for record in _paper_records(papers):
        lines.extend(
            [
                "TY  - JOUR",
                f"TI  - {record['title']}",
            ]
        )
        for author in record["authors"].split("; "):
            if author:
                lines.append(f"AU  - {author}")
        if record["journal"]:
            lines.append(f"JO  - {record['journal']}")
        if record["year"]:
            lines.append(f"PY  - {record['year']}")
        if record["doi"]:
            lines.append(f"DO  - {record['doi']}")
        if record["pmid"]:
            lines.append(f"AN  - {record['pmid']}")
        lines.extend(["ER  - ", ""])
    return "\n".join(lines).strip() + "\n"


def research_result_bibtex(
    question: str, result: dict[str, Any], papers: Sequence[dict[str, Any]]
) -> str:
    """Render the papers backing a research run as BibTeX entries."""
    lines = ["% Science Buddy BibTeX export", f"% Research question: {question}", ""]
    for record in _paper_records(papers):
        key = _bibtex_key(record)
        lines.append(f"@article{{{key},")
        if record["authors"]:
            lines.append(f"  author = {{{record['authors']}}},")
        if record["title"]:
            lines.append(f"  title = {{{{{record['title']}}}}},")
        if record["journal"]:
            lines.append(f"  journal = {{{record['journal']}}},")
        if record["year"]:
            lines.append(f"  year = {{{record['year']}}},")
        if record["doi"]:
            lines.append(f"  doi = {{{record['doi']}}},")
        if record["pmid"]:
            lines.append(f"  pmid = {{{record['pmid']}}},")
        lines.append("}")
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def synthesis_docx(
    title: str,
    manuscript: str,
    *,
    figures: list[tuple[Path, str]] | None = None,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.08)
    section.bottom_margin = Inches(0.98)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.98)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    heading_styles = (
        ("Title", 22),
        ("Heading 1", 16),
        ("Heading 2", 14),
        ("Heading 3", 12),
    )
    for style_name, size in heading_styles:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
    heading_seen = False
    for raw_line in manuscript.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not heading_seen:
                paragraph = document.add_heading(text or title, level=0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                heading_seen = True
            else:
                document.add_heading(text, level=min(level, 3))
            continue
        if line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\.\s+", line):
            document.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
        elif line.startswith(">"):
            document.add_paragraph(line.lstrip("> "), style="Quote")
        elif line.startswith("```") or line.startswith("!["):
            continue
        else:
            document.add_paragraph(line)
    if figures:
        document.add_heading("图表附录", level=1)
        for path, caption in figures:
            if not path.is_file():
                continue
            document.add_picture(str(path), width=Inches(6.2))
            paragraph = document.add_paragraph(caption or path.name)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    output = BytesIO()
    document.save(output)
    return output.getvalue()
