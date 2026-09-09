from collections.abc import Coroutine
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import httpx
import pytest
from docx import Document
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from science_buddy.api import synthesis as synthesis_api
from science_buddy.api.schemas import SynthesisSessionCreateRequest
from science_buddy.config import Settings, get_settings
from science_buddy.domain.enums import JobStatus
from science_buddy.infrastructure.models import (
    Base,
    Job,
    Project,
    SynthesisSection,
    SynthesisSession,
    SynthesisSource,
)
from science_buddy.main import app
from science_buddy.services.background_tasks import recover_interrupted_background_jobs
from science_buddy.services.exports import synthesis_docx, synthesis_markdown
from science_buddy.services.memory_context import (
    MemoryContextItem,
    MemoryContextResult,
)
from science_buddy.services.synthesis import (
    SynthesisOrchestrator,
    context_char_budget,
    output_token_budget,
)
from science_buddy.services.synthesis_template import distilled_thesis_profile
from science_buddy.services.synthesis_types import FigurePlan


class SynthesisTestModel:
    name = "synthesis-test-model"

    def __init__(self) -> None:
        self.structured_requests: list[Any] = []
        self.text_requests: list[Any] = []

    async def generate_structured(self, request):  # type: ignore[no-untyped-def]
        self.structured_requests.append(request)
        if request.operation == "synthesis.document_understanding":
            return {
                "working_title": "增量论文",
                "central_question": "如何保持长文一致性？",
                "thesis_statement": "章节状态与摘要链可以保持长文一致性。",
                "global_summary": "初始全局摘要。",
                "glossary": {"章节状态": "持久化章节正文及摘要"},
                "source_relations": [],
                "outline": [
                    {
                        "section_key": key,
                        "title": title,
                        "purpose": purpose,
                        "retrieval_queries": [],
                    }
                    for key, title, purpose in (
                        ("introduction", "第1章 绪论", "界定问题"),
                        ("analysis", "第2章 分析", "形成论证"),
                        ("conclusion", "第3章 结论", "综合结论"),
                    )
                ],
                "literature_queries": [],
                "figure_plan": [],
            }
        if request.operation.startswith("synthesis.reviewer."):
            return {"verdict": "approved", "concise_summary": "章节逻辑完整。"}
        if request.operation.startswith("synthesis.section_digest."):
            payload = request.user_content
            return {"summary": f"章节摘要：{payload[-100:]}", "evidence_ids": [], "source_ids": []}
        if request.operation.startswith("synthesis.global_review."):
            return {"verdict": "approved", "overall_assessment": "全文结构完整。"}
        raise AssertionError(f"Unexpected operation: {request.operation}")

    async def generate_text(self, request):  # type: ignore[no-untyped-def]
        self.text_requests.append(request)
        payload = request.user_content
        marker = "保留人工修订" if "人工修订" in payload else "模型初稿"
        return f"## {request.operation.rsplit('.', 1)[-1]}\n\n{marker}，形成完整论证。"

    async def stream_text(self, _system: str, _messages):  # type: ignore[no-untyped-def]
        yield ""


class NoRetrievalSynthesisOrchestrator(SynthesisOrchestrator):
    async def _retrieve(self, _synthesis, _query):  # type: ignore[no-untyped-def]
        return []


class CapturingTaskManager:
    def __init__(self) -> None:
        self.coroutines: list[Coroutine[Any, Any, None]] = []

    def start(self, coroutine: Coroutine[Any, Any, None]) -> None:
        self.coroutines.append(coroutine)

    def close(self) -> None:
        for coroutine in self.coroutines:
            coroutine.close()


async def make_database(path: Path) -> tuple[object, async_sessionmaker]:
    engine = create_async_engine(f"sqlite+aiosqlite:///{path}")
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    return engine, sessions


def test_distilled_template_contains_only_aggregate_structure() -> None:
    profile = distilled_thesis_profile()

    assert profile["files_found"] == 19
    assert profile["unique_documents"] == 16
    assert profile["duplicates_removed"] == 3
    assert profile["dominant_page"]["paper"] == "A4"
    assert profile["high_confidence_order"]
    assert "source_root_hash" in profile


def test_synthesis_exports_are_readable() -> None:
    manuscript = "# 示例论文\n\n## 引言\n\n这是可追溯的正文。\n\n- 第一项"

    markdown = synthesis_markdown("示例论文", manuscript).decode("utf-8")
    document = Document(BytesIO(synthesis_docx("示例论文", manuscript)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert markdown.startswith("# 示例论文")
    assert "这是可追溯的正文" in text
    assert document.sections[0].page_width.inches == pytest.approx(8.27, abs=0.01)


def test_synthesis_accepts_one_million_context_and_scales_selection_budget() -> None:
    payload = SynthesisSessionCreateRequest(
        project_id=uuid4(),
        topic="Long-context synthesis",
        max_context_tokens=1_000_000,
    )

    assert payload.max_context_tokens == 1_000_000
    assert context_char_budget(1_000_000, allocation=0.55, cap=550_000) == 550_000
    assert context_char_budget(32_768, allocation=0.55, cap=550_000) < 32_768
    assert output_token_budget(32_768, cap=49_152) == 8192
    assert output_token_budget(1_000_000, cap=49_152) == 49_152


@pytest.mark.asyncio
async def test_synthesis_writer_prioritizes_verified_memory_but_keeps_it_out_of_evidence(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    engine, sessions = await make_database(tmp_path / "synthesis-memory.db")
    now = datetime.now(UTC)
    fact_item = MemoryContextItem(
        memory_type="project_fact",
        entity_id=uuid4(),
        role="confirmed_project_fact",
        text="已确认事实：BRAF 通路与耐药相关。",
        category="verified_research_claim",
        status="active",
        confidence=0.95,
        importance=0.95,
        fused_score=0.9,
        routes=("semantic",),
        source={
            "historical_evidence_ids": ["ev1.old-fact"],
            "evidence_eligible": False,
        },
        related=(),
        created_at=now,
    )
    claim_item = MemoryContextItem(
        memory_type="verified_claim",
        entity_id=uuid4(),
        role="verified_claim_prior",
        text="已验证 Claim：抑制 BRAF 可改变耐药表型。",
        category="verified_evidence_claim",
        status="supported",
        confidence=0.93,
        importance=1.0,
        fused_score=0.88,
        routes=("semantic",),
        source={
            "historical_evidence_ids": ["ev1.old-claim"],
            "evidence_eligible": False,
            "citation_requires_current_retrieval": True,
        },
        related=(),
        created_at=now,
    )
    step_item = MemoryContextItem(
        memory_type="research_step",
        entity_id=uuid4(),
        role="planning_experience",
        text="历史搜索经验：先使用机制关键词再扩展同义词。",
        category="retrieve",
        status="completed",
        confidence=1.0,
        importance=0.7,
        fused_score=0.7,
        routes=("keyword",),
        source={"historical_evidence_count": 2, "evidence_eligible": False},
        related=(),
        created_at=now,
    )
    context = MemoryContextResult(
        consumer="synthesis",
        allowed_memory_types=(
            "project_fact",
            "verified_claim",
            "research_step",
            "failed_route",
        ),
        items=(fact_item, claim_item, step_item),
        indexed=3,
        embedding_model="test",
    )

    async def recall_context(*_args: object, **_kwargs: object) -> MemoryContextResult:
        return context

    monkeypatch.setattr(
        "science_buddy.services.synthesis.MemoryContextService.recall",
        recall_context,
    )
    async with sessions() as session:
        project = Project(name="Synthesis memory project")
        session.add(project)
        await session.flush()
        synthesis = SynthesisSession(
            project_id=project.id,
            title="Memory thesis",
            topic="BRAF resistance",
            include_workbench_notes=False,
            global_summary="全局摘要",
            document_map={"thesis_statement": "测试长期记忆边界"},
        )
        session.add(synthesis)
        await session.flush()
        section = SynthesisSection(
            session_id=synthesis.id,
            section_key="mechanism",
            ordinal=0,
            level=1,
            title="机制分析",
            purpose="解释耐药机制",
            status="pending",
        )
        session.add(section)
        await session.flush()
        model = SynthesisTestModel()
        orchestrator = NoRetrievalSynthesisOrchestrator(
            session,
            model=model,
            settings=Settings(_env_file=None, upload_directory=tmp_path / "uploads"),
            session_factory=sessions,
        )
        await orchestrator._write_section(synthesis, section, [])  # noqa: SLF001

    writer_request = next(
        value
        for value in model.text_requests
        if value.operation.startswith("synthesis.writer.")
    )
    writer_payload = writer_request.user_content
    assert "verified_prior_context" in writer_payload
    assert "已确认事实" in writer_payload
    assert "已验证 Claim" in writer_payload
    assert "planning_experience_not_evidence" in writer_payload
    assert "历史搜索经验" in writer_payload
    assert '"evidence": []' in writer_payload
    assert "historical Evidence IDs are not valid citations" in writer_request.system_instruction
    await engine.dispose()  # type: ignore[attr-defined]


@pytest.mark.asyncio
async def test_synthesis_api_creates_session_deduplicates_sources_and_exports(
    tmp_path: Path,
) -> None:
    engine, sessions = await make_database(tmp_path / "synthesis-api.db")
    settings = Settings(_env_file=None, upload_directory=tmp_path / "uploads")
    async with sessions() as session:
        project = Project(name="Synthesis project")
        session.add(project)
        await session.commit()
        project_id = project.id

    async def override_session():  # type: ignore[no-untyped-def]
        async with sessions() as session:
            yield session

    def override_settings() -> Settings:
        return settings

    app.dependency_overrides[synthesis_api.get_session] = override_session
    app.dependency_overrides[get_settings] = override_settings
    transport = httpx.ASGITransport(app=app)
    try:
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            created = await client.post(
                "/api/v1/synthesis/sessions",
                json={
                    "project_id": str(project_id),
                    "topic": "整合工作台笔记并形成完整论文",
                    "model_depth": "deep",
                    "max_context_tokens": 65536,
                    "max_review_rounds": 2,
                    "include_workbench_notes": True,
                    "allow_online_literature": False,
                },
            )
            assert created.status_code == 201, created.text
            synthesis_id = UUID(created.json()["id"])
            first = await client.post(
                f"/api/v1/synthesis/sessions/{synthesis_id}/sources/text",
                json={"title": "实验说明", "content": "# 结果\n\n观察到稳定趋势。"},
            )
            duplicate = await client.post(
                f"/api/v1/synthesis/sessions/{synthesis_id}/sources/text",
                json={"title": "重复说明", "content": "# 结果\n\n观察到稳定趋势。"},
            )
            assert first.status_code == 200, first.text
            assert duplicate.status_code == 200, duplicate.text
            assert duplicate.json()["id"] == first.json()["id"]
            uploaded = await client.post(
                f"/api/v1/synthesis/sessions/{synthesis_id}/sources/upload",
                files={
                    "file": (
                        "additional.md",
                        b"# Additional evidence\n\nA separately uploaded source.",
                        "text/markdown",
                    )
                },
            )
            assert uploaded.status_code == 200, uploaded.text
            assert uploaded.json()["filename"] == "additional.md"

            async with sessions() as session:
                synthesis = await session.get(SynthesisSession, synthesis_id)
                assert synthesis is not None
                synthesis.manuscript_markdown = "# 完整论文\n\n## 引言\n\n正文内容。"
                section = SynthesisSection(
                    session_id=synthesis_id,
                    section_key="introduction",
                    ordinal=0,
                    level=1,
                    title="引言",
                    purpose="界定问题",
                    draft_markdown="## 引言\n\n正文内容。",
                    status="reviewed",
                )
                session.add(section)
                await session.commit()
                section_id = section.id

            detail = await client.get(f"/api/v1/synthesis/sessions/{synthesis_id}")
            markdown = await client.get(
                f"/api/v1/synthesis/sessions/{synthesis_id}/export?format=markdown"
            )
            docx = await client.get(
                f"/api/v1/synthesis/sessions/{synthesis_id}/export?format=docx"
            )
            updated = await client.patch(
                f"/api/v1/synthesis/sessions/{synthesis_id}/sections/{section_id}",
                json={"draft_markdown": "## 引言\n\n人工修订后的正文。"},
            )
    finally:
        app.dependency_overrides.clear()

    assert detail.status_code == 200, detail.text
    assert detail.json()["source_count"] == 2
    assert detail.json()["section_count"] == 1
    assert markdown.status_code == 200
    assert markdown.content.startswith(b"# ")
    assert docx.status_code == 200
    assert Document(BytesIO(docx.content)).paragraphs
    assert updated.status_code == 200, updated.text
    assert updated.json()["revision"] == 1
    assert updated.json()["status"] == "user_edited"
    await engine.dispose()  # type: ignore[attr-defined]


@pytest.mark.asyncio
async def test_synthesis_run_endpoint_reuses_active_job(tmp_path: Path) -> None:
    engine, sessions = await make_database(tmp_path / "synthesis-job.db")
    settings = Settings(_env_file=None, upload_directory=tmp_path / "uploads")
    async with sessions() as session:
        project = Project(name="Background project")
        session.add(project)
        await session.flush()
        synthesis = SynthesisSession(
            project_id=project.id,
            title="Background thesis",
            topic="Persistent jobs",
        )
        session.add(synthesis)
        await session.commit()
        synthesis_id = synthesis.id

    async def override_session():  # type: ignore[no-untyped-def]
        async with sessions() as session:
            yield session

    def override_settings() -> Settings:
        return settings

    manager = CapturingTaskManager()
    app.state.background_tasks = manager
    app.dependency_overrides[synthesis_api.get_session] = override_session
    app.dependency_overrides[get_settings] = override_settings
    transport = httpx.ASGITransport(app=app)
    try:
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            first = await client.post(f"/api/v1/synthesis/sessions/{synthesis_id}/run")
            duplicate = await client.post(f"/api/v1/synthesis/sessions/{synthesis_id}/run")
    finally:
        app.dependency_overrides.clear()
        manager.close()

    assert first.status_code == 202, first.text
    assert duplicate.status_code == 202, duplicate.text
    assert duplicate.json()["id"] == first.json()["id"]
    assert first.json()["status"] == "queued"
    assert len(manager.coroutines) == 1
    await engine.dispose()  # type: ignore[attr-defined]


@pytest.mark.asyncio
async def test_interrupted_synthesis_job_is_failed_and_session_unlocked(
    tmp_path: Path,
) -> None:
    engine, sessions = await make_database(tmp_path / "synthesis-recovery.db")
    async with sessions() as session:
        project = Project(name="Recovery project")
        session.add(project)
        await session.flush()
        synthesis = SynthesisSession(
            project_id=project.id,
            title="Interrupted thesis",
            topic="Recovery test",
            status="writing",
        )
        session.add(synthesis)
        await session.flush()
        job = Job(
            kind="synthesis_manuscript",
            status=JobStatus.RUNNING,
            idempotency_key="synthesis-recovery-test",
            payload={"session_id": str(synthesis.id)},
        )
        session.add(job)
        await session.commit()
        synthesis_id = synthesis.id
        job_id = job.id

    recovered = await recover_interrupted_background_jobs(sessions)

    assert recovered == 1
    async with sessions() as session:
        synthesis = await session.get(SynthesisSession, synthesis_id)
        job = await session.get(Job, job_id)
        assert synthesis is not None
        assert job is not None
        assert synthesis.status == "failed"
        assert "重新提交" in (synthesis.error_message or "")
        assert job.status == JobStatus.FAILED
    await engine.dispose()  # type: ignore[attr-defined]


@pytest.mark.asyncio
async def test_synthesis_plot_agent_renders_real_markdown_table(tmp_path: Path) -> None:
    engine, sessions = await make_database(tmp_path / "synthesis-plot.db")
    settings = Settings(_env_file=None, upload_directory=tmp_path / "uploads")
    async with sessions() as session:
        project = Project(name="Plot project")
        session.add(project)
        await session.flush()
        synthesis = SynthesisSession(
            project_id=project.id,
            title="Plot thesis",
            topic="Plot real data",
            include_workbench_notes=False,
        )
        session.add(synthesis)
        await session.flush()
        source = SynthesisSource(
            session_id=synthesis.id,
            source_type="user_text",
            filename="replicates.md",
            media_type="text/markdown",
            content_hash="a" * 64,
            extracted_text=(
                "| 条件 | 重复1 | 重复2 | 重复3 |\n"
                "| --- | ---: | ---: | ---: |\n"
                "| 对照 | 1.0 | 1.2 | 0.9 |\n"
                "| 处理 | 2.1 | 2.4 | 2.2 |"
            ),
            status="ready",
        )
        session.add(source)
        await session.commit()
        orchestrator = NoRetrievalSynthesisOrchestrator(
            session,
            model=SynthesisTestModel(),
            settings=settings,
            session_factory=sessions,
        )
        await orchestrator._generate_figures(  # noqa: SLF001
            synthesis,
            [source],
            [
                FigurePlan(
                    source_id=str(source.id),
                    table_index=0,
                    title="处理组重复实验",
                    caption="柱形、误差与全部重复点。",
                )
            ],
        )
        await session.refresh(synthesis)
        manifest: list[dict[str, Any]] = synthesis.figure_manifest

    assert len(manifest) == 1
    assert manifest[0]["chart_type"] == "bar"
    assert manifest[0]["source_id"] == str(source.id)
    assert Path(manifest[0]["storage_path"]).is_file()
    assert manifest[0]["render_spec"]["data_layout"] == "wide"
    assert len(manifest[0]["render_spec"]["replicate_stats"]) == 2
    await engine.dispose()  # type: ignore[attr-defined]


@pytest.mark.asyncio
async def test_synthesis_rerun_preserves_section_identity_and_iterates_draft(
    tmp_path: Path,
) -> None:
    engine, sessions = await make_database(tmp_path / "synthesis-iteration.db")
    settings = Settings(_env_file=None, upload_directory=tmp_path / "uploads")
    async with sessions() as session:
        project = Project(name="Iteration project")
        session.add(project)
        await session.flush()
        synthesis = SynthesisSession(
            project_id=project.id,
            title="Initial title",
            topic="Long manuscript consistency",
            include_workbench_notes=False,
            allow_online_literature=False,
            max_review_rounds=1,
        )
        session.add(synthesis)
        await session.commit()
        orchestrator = NoRetrievalSynthesisOrchestrator(
            session,
            model=SynthesisTestModel(),
            settings=settings,
            session_factory=sessions,
        )
        await orchestrator.run(synthesis)
        first_sections = await orchestrator._sections(synthesis.id)  # noqa: SLF001
        original_ids = [value.id for value in first_sections]
        first_sections[0].draft_markdown = "## 第1章 绪论\n\n人工修订必须保留。"
        first_sections[0].status = "user_edited"
        await session.commit()

        await orchestrator.run(synthesis)
        second_sections = await orchestrator._sections(synthesis.id)  # noqa: SLF001
        await session.refresh(synthesis)

    assert [value.id for value in second_sections] == original_ids
    assert "保留人工修订" in second_sections[0].draft_markdown
    assert all(value.revision == 2 for value in second_sections)
    assert set(synthesis.citation_ledger) == {"introduction", "analysis", "conclusion"}
    assert "[introduction]" in synthesis.global_summary
    assert synthesis.current_round == 2
    await engine.dispose()  # type: ignore[attr-defined]
